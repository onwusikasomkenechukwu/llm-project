#!/usr/bin/env python3
"""Preflight a benchmark run before spending API budget.

    python check_run.py prompts.xlsx
    python check_run.py prompts.xlsx --expected-questions 400 --polarity pos-neg
    python check_run.py prompts.xlsx --responses responses.jsonl --judgments judgments.jsonl

This checks the spreadsheet contract, computes expected row counts, and can
audit merged response and judgment files against the run grid. It calls no APIs.
"""

import argparse
import json
import os
import sys
import subprocess
from zipfile import BadZipFile
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


IDENTITIES = (
    "none",
    "black-american",
    "white-american",
    "latino-american",
    "christian",
    "muslim",
    "jewish",
    "conservative",
    "progressive",
    "libertarian",
    "woman",
    "man",
    "learning-disability",
    "disabled-veteran",
    "immigrant",
    "native-american",
    "transgender-woman",
    "non-binary",
)

PROVIDERS = ("openai", "anthropic", "google", "xai", "muse")
QUESTION_TYPES = (
    "Factual",
    "Directed",
    "Loaded (False Premise)",
    "Loaded (True Premise)",
    "Open-Ended",
)
CRITERIA = ("accuracy", "completeness", "objectivity", "sourcing")


class Report:
    def __init__(self) -> None:
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def error(self, text: str) -> None:
        self.errors.append(text)

    def warn(self, text: str) -> None:
        self.warnings.append(text)

    def print(self) -> None:
        for text in self.errors:
            print(f"[ERROR] {text}")
        for text in self.warnings:
            print(f"[WARN]  {text}")
        if not self.errors and not self.warnings:
            print("[OK] no preflight issues found")
        else:
            print(f"[SUMMARY] {len(self.errors)} error(s), {len(self.warnings)} warning(s)")


def split_csv(raw: str, known: tuple[str, ...], label: str, report: Report) -> tuple[str, ...]:
    items = tuple(x.strip() for x in raw.split(",") if x.strip())
    if not items:
        report.error(f"empty {label} selection")
    unknown = [x for x in items if x not in known]
    if unknown:
        report.error(f"unknown {label}: {unknown}. Known: {list(known)}")
    if len(set(items)) != len(items):
        report.error(f"duplicate {label} in option: {raw}")
    return items


def pick_column(header: list[str], explicit: str | None, guesses: tuple[str, ...],
                label: str, report: Report) -> int | None:
    lower = {h.lower(): i for i, h in enumerate(header) if h}
    if explicit:
        j = lower.get(explicit.lower())
        if j is None:
            report.error(f"no column named {explicit!r} for {label}. Columns: {header}")
        return j
    return next((lower[g] for g in guesses if g in lower), None)


def cell(row: tuple[object, ...], j: int | None) -> str:
    if j is None or j >= len(row) or row[j] is None:
        return ""
    return str(row[j]).strip()


def load_prompts(args: argparse.Namespace, report: Report) -> tuple[list[dict], list[dict]]:
    path = Path(args.prompts)
    if not path.exists():
        report.error(f"prompt workbook not found: {path}")
        return [], []

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except (OSError, ValueError, BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as exc:
        report.error(f"cannot read prompt workbook: {exc}")
        return [], []
    sheet = int(args.sheet) if str(args.sheet).isdigit() else args.sheet
    if isinstance(sheet, str):
        if sheet not in wb.sheetnames:
            report.error(f"sheet {sheet!r} not found. Sheets: {wb.sheetnames}")
            wb.close()
            return [], []
        ws = wb[sheet]
    else:
        if sheet < 0 or sheet >= len(wb.worksheets):
            report.error(f"sheet index {sheet} out of range. Sheets: {wb.sheetnames}")
            wb.close()
            return [], []
        ws = wb.worksheets[sheet]

    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        report.error(f"{path} sheet {ws.title!r} is empty")
        return [], []

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    names = [h.lower() for h in header if h]
    if len(names) != len(set(names)):
        report.error("duplicate column names in prompt workbook")
    pos = pick_column(header, args.pos_col,
                      ("question", "base prompt", "prompt", "claim", "positive",
                       "statement", "original"),
                      "question", report)
    neg = pick_column(header, args.neg_col,
                      ("negative", "negation", "negated", "opposite", "reversed"),
                      "negation", report)
    qid_col = pick_column(header, args.id_col,
                          ("prompt id", "id", "qid", "question_id", "item", "index",
                           "no", "number"),
                          "question id", report)
    qtype_col = pick_column(header, None,
                            ("question type", "prompt type", "type", "structure"),
                            "question type", report)
    ideal_col = pick_column(header, None,
                            ("ideal answer", "reference answer", "ideal",
                             "expected answer", "model answer"),
                            "reference answer", report)
    components_col = pick_column(header, None,
                                 ("ideal answer components", "components", "key points",
                                  "must cover", "required elements"),
                                 "reference components", report)

    if pos is None:
        report.error(f"could not find the question column. Pass --pos-col. Columns: {header}")
        return [], []
    if neg is None and args.polarity == "pos-neg":
        report.error("requested --polarity pos-neg, but no negation column was found")
    elif neg is None and args.polarity == "auto":
        report.warn("no negation column found, so only positive questions are expected")
    if qtype_col is None:
        report.warn("no question-type column found, so premise-handling analysis is weakened")

    questions: list[dict] = []
    cells: list[dict] = []
    seen_qids: Counter[str] = Counter()
    missing_ids = 0

    for row_n, row in enumerate(rows[1:], start=2):
        qid = cell(row, qid_col)
        if not qid:
            missing_ids += 1
            qid = f"q{row_n - 1:04d}"
        qtype = cell(row, qtype_col) or "Unspecified"
        positive = cell(row, pos)
        negative = cell(row, neg)
        ideal = cell(row, ideal_col)
        components = cell(row, components_col)
        if not positive and not negative:
            report.warn(f"row {row_n} has neither a positive nor negative question")
            continue
        if "|" in qid:
            report.error(f"row {row_n} qid {qid!r} contains '|', which breaks output ids")
        if qtype not in QUESTION_TYPES:
            report.warn(f"row {row_n} qid {qid!r} has nonstandard question type {qtype!r}")
        if args.require_references and not (ideal or components):
            report.error(f"row {row_n} qid {qid!r} has no reference answer or components")
        if args.require_question_types and qtype not in QUESTION_TYPES:
            report.error(f"row {row_n} qid {qid!r} has no recognized question type")
        seen_qids[qid] += 1
        questions.append({
            "qid": qid,
            "row": row_n,
            "qtype": qtype,
            "has_reference": bool(ideal or components),
            "has_pos": bool(positive),
            "has_neg": bool(negative),
        })

        if args.polarity in ("auto", "pos", "pos-neg") and positive:
            cells.append({"qid": qid, "polarity": "pos", "qtype": qtype})
        if args.polarity in ("auto", "pos-neg") and negative:
            cells.append({"qid": qid, "polarity": "neg", "qtype": qtype})

    if args.limit is not None and args.limit > 0:
        keep = {q["qid"] for q in questions[:args.limit]}
        questions = [q for q in questions if q["qid"] in keep]
        cells = [c for c in cells if c["qid"] in keep]

    for qid, n in seen_qids.items():
        if n > 1:
            report.error(f"duplicate qid {qid!r} appears {n} times")
    if missing_ids:
        report.warn(f"{missing_ids} row(s) have no qid and will fall back to q0001-style ids")
    if args.expected_questions is not None and len(questions) != args.expected_questions:
        report.error(f"expected {args.expected_questions} questions, found {len(questions)}")
    if args.polarity == "pos-neg":
        missing_pos = [q["qid"] for q in questions if not q["has_pos"]]
        if missing_pos:
            report.error(f"{len(missing_pos)} question(s) are missing positive questions")
        missing_neg = [q["qid"] for q in questions if not q["has_neg"]]
        if missing_neg:
            report.error(f"{len(missing_neg)} question(s) are missing negations, e.g. {missing_neg[:5]}")

    if not cells:
        report.error("no usable prompt cells for the selected polarity")
    return questions, cells


def expected_answer_ids(cells: list[dict], providers: tuple[str, ...],
                        identities: tuple[str, ...], duplicates: int) -> set[str]:
    ids = set()
    for q in cells:
        for identity in identities:
            for provider in providers:
                for run in range(duplicates):
                    ids.add(f"{q['qid']}|{q['polarity']}|{identity}|{provider}|{run}")
    return ids


def read_jsonl(path: Path, report: Report) -> list[dict]:
    rows = []
    if not path.is_file():
        report.error(f"jsonl file not found: {path}")
        return rows
    with path.open(encoding="utf-8") as f:
        for n, line in enumerate(f, 1):
            line = line.strip()
            if not line:
                continue
            try:
                row = json.loads(line)
                if not isinstance(row, dict):
                    report.error(f"{path}:{n} must contain a JSON object")
                else:
                    rows.append(row)
            except json.JSONDecodeError as exc:
                report.error(f"{path}:{n} is not valid JSON: {exc.msg}")
    return rows


def id_parts(row_id: object) -> tuple[str, str, str, str, str] | None:
    if not isinstance(row_id, str):
        return None
    parts = row_id.split("|")
    if len(parts) != 5:
        return None
    return tuple(parts)  # type: ignore[return-value]


def audit_responses(path: Path, expected_ids: set[str], providers: tuple[str, ...],
                    identities: tuple[str, ...], report: Report) -> None:
    rows = read_jsonl(path, report)

    by_id: defaultdict[str, list[dict]] = defaultdict(list)
    errors = 0
    success = 0
    unexpected: list[str] = []

    for row in rows:
        row_id = row.get("id")
        by_id[str(row_id)].append(row)
        parts = id_parts(row_id)
        if parts is None:
            report.error(f"response row has malformed id: {row_id!r}")
            continue
        qid, polarity, identity, provider, run = parts
        if row_id not in expected_ids:
            unexpected.append(str(row_id))
        if row.get("qid") not in (None, qid):
            report.error(f"response {row_id} has qid field {row.get('qid')!r}")
        if row.get("polarity") not in (None, polarity):
            report.error(f"response {row_id} has polarity field {row.get('polarity')!r}")
        for name, value in (("identity", identity), ("provider", provider)):
            if row.get(name) != value:
                report.error(f"response {row_id} has inconsistent {name}")
        if type(row.get("run")) is not int or str(row["run"]) != run:
            report.error(f"response {row_id} has inconsistent run")
        if identity not in identities:
            report.error(f"response {row_id} has unknown identity {identity!r}")
        if provider not in providers:
            report.error(f"response {row_id} has unknown provider {provider!r}")
        try:
            if int(run) < 0:
                report.error(f"response {row_id} has negative run")
        except ValueError:
            report.error(f"response {row_id} has non-integer run {run!r}")
        if row.get("error") is None:
            success += 1
            if not isinstance(row.get("response"), str) or not row["response"].strip():
                report.error(f"response {row_id} succeeded but has empty response text")
        else:
            errors += 1

    duplicates = {k: len(v) for k, v in by_id.items() if len(v) > 1}
    if duplicates:
        report.error(f"{len(duplicates)} response id(s) appear more than once; merge first")
    if unexpected:
        report.error(f"{len(unexpected)} response row(s) are outside the expected grid, e.g. {unexpected[:5]}")
    missing = expected_ids - set(by_id)
    if missing:
        report.error(f"{len(missing)} expected response row(s) are missing, e.g. {sorted(missing)[:5]}")
    if errors:
        report.error(f"{errors} response row(s) still have errors")

    print(f"responses: {len(rows)} row(s), {success} success, {errors} with error")


def valid_score(value: object) -> bool:
    return type(value) is int and 1 <= value <= 5


def audit_judgments(path: Path, expected_answer_ids_: set[str], judges: tuple[str, ...],
                    passes: int, report: Report) -> None:
    rows = read_jsonl(path, report)

    expected_keys = {
        (answer_id, judge, pass_n)
        for answer_id in expected_answer_ids_
        for judge in judges
        for pass_n in range(passes)
    }
    seen_keys: set[tuple[str, str, int]] = set()
    errors = 0
    success = 0
    duplicates = 0

    for row in rows:
        row_id = row.get("id")
        judge = row.get("judge")
        pass_n = row.get("pass", 0)
        if not isinstance(row_id, str) or id_parts(row_id) is None:
            report.error(f"judgment row has malformed answer id: {row_id!r}")
            continue
        if judge not in judges:
            report.error(f"judgment {row_id} has unknown judge {judge!r}")
        if type(pass_n) is not int or pass_n < 0:
            report.error(f"judgment {row_id} has invalid pass {pass_n!r}")
            pass_n = -1
        key = (row_id, str(judge), pass_n)
        if key in seen_keys:
            duplicates += 1
        seen_keys.add(key)
        parts = id_parts(row_id)
        for name, value in zip(("qid", "polarity", "identity", "provider"), parts[:4]):
            if row.get(name) != value:
                report.error(f"judgment {row_id} has inconsistent {name}")

        if row.get("error") is not None:
            errors += 1
            continue
        success += 1
        missing_scores = [k for k in CRITERIA if not valid_score(row.get(k))]
        if missing_scores:
            report.error(f"judgment {row_id} by {judge} has invalid scores: {missing_scores}")
        elif row.get("total") != sum(int(row[k]) for k in CRITERIA):
            report.error(f"judgment {row_id} by {judge} total does not match component scores")
        if not isinstance(row.get("hard_fail"), bool):
            report.error(f"judgment {row_id} by {judge} has non-boolean hard_fail")
        qtype = row.get("question_type") or "Unspecified"
        loaded = str(qtype).startswith("Loaded")
        premise = row.get("premise_handling")
        if loaded and not valid_score(premise):
            report.error(f"judgment {row_id} by {judge} is loaded but premise_handling is missing")
        if not loaded and premise is not None:
            report.error(f"judgment {row_id} by {judge} is not loaded but has premise_handling")
        if not row.get("raw"):
            report.warn(f"judgment {row_id} by {judge} has no raw grader reply")

    if duplicates:
        report.error(f"{duplicates} duplicate judgment key(s); merge first")
    unexpected = seen_keys - expected_keys
    if unexpected:
        sample = sorted(unexpected)[:5]
        report.error(f"{len(unexpected)} judgment key(s) are outside the expected grid, e.g. {sample}")
    missing = expected_keys - seen_keys
    if missing:
        sample = sorted(missing)[:5]
        report.error(f"{len(missing)} expected judgment key(s) are missing, e.g. {sample}")
    if errors:
        report.error(f"{errors} judgment row(s) still have errors")

    print(f"judgments: {len(rows)} row(s), {success} success, {errors} with error")


def print_prompt_summary(questions: list[dict], cells: list[dict],
                         providers: tuple[str, ...], judges: tuple[str, ...],
                         identities: tuple[str, ...], duplicates: int,
                         passes: int) -> None:
    by_type = Counter(q["qtype"] for q in questions)
    refs = sum(1 for q in questions if q["has_reference"])
    positives = sum(1 for q in questions if q["has_pos"])
    negatives = sum(1 for q in questions if q["has_neg"])
    answers = len(cells) * len(identities) * len(providers) * duplicates
    judgments = answers * len(judges) * passes

    print(f"questions: {len(questions)} total, {positives} positive, {negatives} negated")
    print(f"references: {refs} question(s) with reference answer or components")
    if by_type:
        print("question types: " + ", ".join(f"{k}={v}" for k, v in sorted(by_type.items())))
    print(f"grid: {len(identities)} identities, {len(providers)} answerer(s), "
          f"{len(judges)} judge(s), {duplicates} duplicate(s), {passes} pass(es)")
    print(f"expected answers: {answers}")
    print(f"expected judgments: {judgments}")


def verify_drift() -> None:
    guard = os.path.join(os.path.dirname(os.path.abspath(__file__)), "check_drift.py")
    result = subprocess.run([sys.executable, guard], capture_output=True, text=True)
    if result.returncode:
        sys.exit(result.stderr.strip() or result.stdout.strip() or "Drift guard failed.")


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("prompts", help="prompts .xlsx to check")
    p.add_argument("--responses", help="optional responses.jsonl to audit")
    p.add_argument("--judgments", help="optional judgments.jsonl to audit")
    p.add_argument("--expected-questions", type=int)
    p.add_argument("--polarity", choices=["auto", "pos", "pos-neg"], default="auto",
                   help="expected prompt polarities. auto follows non-empty spreadsheet cells")
    p.add_argument("--providers", default=",".join(PROVIDERS))
    p.add_argument("--judges", default=",".join(PROVIDERS))
    p.add_argument("--identities", default=",".join(IDENTITIES))
    p.add_argument("--duplicates", type=int, default=3)
    p.add_argument("--passes", type=int, default=1)
    p.add_argument("--limit", type=int, help="check the first N questions, like the run scripts")
    p.add_argument("--require-references", action="store_true")
    p.add_argument("--require-question-types", action="store_true")
    p.add_argument("--sheet", default=0)
    p.add_argument("--id-col")
    p.add_argument("--pos-col")
    p.add_argument("--neg-col")
    args = p.parse_args()
    verify_drift()

    report = Report()
    providers = split_csv(args.providers, PROVIDERS + ("local",), "provider(s)", report)
    judges = split_csv(args.judges, PROVIDERS + ("local",), "judge(s)", report)
    identities = split_csv(args.identities, IDENTITIES, "identity/identities", report)
    if args.duplicates < 1:
        report.error("--duplicates must be at least 1")
    if args.passes < 1:
        report.error("--passes must be at least 1")
    if args.limit is not None and args.limit < 1:
        report.error("--limit must be at least 1")

    questions, cells = load_prompts(args, report)
    print_prompt_summary(questions, cells, providers, judges, identities,
                         args.duplicates, args.passes)

    expected_ids = expected_answer_ids(cells, providers, identities, args.duplicates)
    if args.responses:
        audit_responses(Path(args.responses), expected_ids, providers, identities, report)
    if args.judgments:
        audit_judgments(Path(args.judgments), expected_ids, judges, args.passes, report)

    report.print()
    return 1 if report.errors else 0


if __name__ == "__main__":
    sys.exit(main())
