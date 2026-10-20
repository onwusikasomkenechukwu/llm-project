#!/usr/bin/env python3
"""Build the prompts workbook from the law school's team sheet.

    python samples/make_prompts.py "Civil Rights and AI Benchmark Project - Team Collaboration Template.xlsx"

Writes samples/prompts-<today>.xlsx: one row per written question across every
topic tab, in the column layout the run scripts expect. A tab counts as a topic
if its header row has a Prompt ID and a Question (or Base Prompt) column; rows
with no question text yet are skipped.

Question types are mapped onto the rubric's five names, since the sheet spells
them several ways ("Loaded True Premise", "Open-Ended/Neutral"). Premise
handling is gated on these names, so a type that maps to none of them stops the
build rather than being graded as the wrong kind.
"""

import datetime
import pathlib
import re
import sys

TYPES = (("false", "Loaded (False Premise)"), ("true", "Loaded (True Premise)"),
         ("open", "Open-Ended"), ("fact", "Factual"), ("direct", "Directed"))


def qtype(raw):
    s = raw.lower()
    return next((name for key, name in TYPES if key in s), None)


def main():
    import openpyxl

    if len(sys.argv) != 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
    out, seen = [], set()
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        head = [re.sub(r"\s+", " ", str(c)).strip().lower() if c else "" for c in rows[0]]
        q = next((head.index(h) for h in ("question", "base prompt") if h in head), None)
        if "prompt id" not in head or q is None:
            continue
        pid, typ = head.index("prompt id"), head.index("question type")
        gold = next((j for j, h in enumerate(head) if h.startswith("gold-standard")), None)

        def cell(r, j):
            return "" if j is None or j >= len(r) or r[j] is None else str(r[j]).strip()

        for r in rows[1:]:
            if not cell(r, q):
                continue
            ident = re.sub(r"\s+", "", cell(r, pid))   # the sheet has "CE - 020"
            kind = qtype(cell(r, typ))
            if not ident or ident in seen:
                sys.exit(f"{ws.title}: missing or duplicate Prompt ID {ident!r}")
            if kind is None:
                sys.exit(f"{ident}: unknown question type {cell(r, typ)!r}")
            seen.add(ident)
            out.append([ident, ws.title, kind, cell(r, q), cell(r, gold)])

    target = pathlib.Path(__file__).with_name(f"prompts-{datetime.date.today()}.xlsx")
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "prompts"
    sheet.append(["Prompt ID", "Topic", "Question Type", "Question", "Ideal Answer Components"])
    for row in out:
        sheet.append(row)
    book.save(target)
    graded = sum(bool(r[4]) for r in out)
    print(f"wrote {target}: {len(out)} questions, {graded} with gold-standard elements")


if __name__ == "__main__":
    main()
