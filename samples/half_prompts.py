#!/usr/bin/env python3
"""Take a balanced half of a prompts workbook, for the demographic cue screen.

    python samples/half_prompts.py samples/prompts-2026-10-20.xlsx

Writes prompts-half.xlsx beside the input: half of every topic x question type
cell, so the half keeps the full set's mix. Within a cell the questions are
ordered by a hash of their id, so the same half comes out every time. A cell
with an odd count gives up its extra question to every other odd cell in turn,
which lands the total on half of the whole, rounded down.

Also writes prompts-probe.xlsx, one placeholder question, for the cue sheet's
probes: they carry no {QUESTION}, so each runs once per model per replicate.
"""

import hashlib
import pathlib
import sys
from collections import defaultdict


def main():
    import openpyxl

    if len(sys.argv) != 2:
        sys.exit(__doc__)
    source = pathlib.Path(sys.argv[1])
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    head, body = list(rows[0]), [r for r in rows[1:] if any(c is not None for c in r)]
    topic, qtype, pid = (head.index(h) for h in ("Topic", "Question Type", "Prompt ID"))

    cells = defaultdict(list)
    for r in body:
        cells[(r[topic], r[qtype])].append(r)
    keep, odd = [], 0
    # type before topic, so the extras alternate across topics within each type
    for key in sorted(cells, key=lambda k: (k[1], k[0])):
        group = sorted(cells[key], key=lambda r: hashlib.sha256(str(r[pid]).encode()).hexdigest())
        n = len(group) // 2
        if len(group) % 2:
            n += odd % 2
            odd += 1
        keep += group[:n]

    def save(path, rows_out):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "prompts"
        sheet.append(head)
        for r in rows_out:
            sheet.append(list(r))
        book.save(path)

    order = {r[pid]: i for i, r in enumerate(body)}
    keep.sort(key=lambda r: order[r[pid]])
    half = source.with_name("prompts-half.xlsx")
    save(half, keep)
    probe = [None] * len(head)
    probe[pid], probe[topic], probe[qtype] = "PROBE", "probe", "Unspecified"
    probe[head.index("Question")] = "(probe: the cue is the whole prompt)"
    save(source.with_name("prompts-probe.xlsx"), [probe])

    print(f"wrote {half}: {len(keep)} of {len(body)} questions")
    tally = defaultdict(int)
    for r in keep:
        tally[r[qtype]] += 1
    print("  by type: " + ", ".join(f"{k} {v}" for k, v in sorted(tally.items())))
    print(f"  topics: {len({r[topic] for r in keep})} of {len({r[topic] for r in body})}")


if __name__ == "__main__":
    main()
